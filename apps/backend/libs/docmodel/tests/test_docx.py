import io
from typing import Any

import docx
import pytest
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from docmodel import BlockKind, DocFormat, SpanKind, export, mask, parse
from docmodel.docx_format import iter_paragraphs


def _hyperlink(paragraph: Any, text: str) -> None:
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), "sec2")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _marker(tag: str, **attrs: str) -> Any:
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn(f"w:{k}"), v)
    return el


def make_docx() -> bytes:
    d = docx.Document()
    d.add_heading("Introduction", level=1)
    p = d.add_paragraph()
    p._p.append(_marker("w:bookmarkStart", id="0", name="intro"))
    p.add_run("Moreover, the ")
    p.add_run("in vitro").italic = True
    p.add_run(" results play a crucial role in this field (Smith, 2020), as shown in ")
    _hyperlink(p, "Section 2")
    p.add_run(". It is important to note that 30% of samples failed.")
    p._p.append(_marker("w:bookmarkEnd", id="0"))
    d.add_paragraph(
        "A bulleted point that is long enough to be rewritten by us.", style="List Bullet"
    )
    table = d.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Cell text that should never be rewritten in the MVP version."
    d.add_heading("References", level=1)
    d.add_paragraph("Smith, J. (2020). A study of things. Journal of Stuff, 4(2), 1-10.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_parse_structure_and_inline_spans() -> None:
    doc = parse(make_docx(), DocFormat.DOCX)
    assert [b.kind for b in doc.blocks] == [
        BlockKind.HEADING,
        BlockKind.PARAGRAPH,
        BlockKind.LIST_ITEM,
        BlockKind.TABLE_CELL,
        BlockKind.HEADING,
        BlockKind.REFERENCE_ENTRY,
    ]
    para = doc.blocks[1]
    assert [s.text for s in para.spans if s.kind is SpanKind.INLINE] == ["in vitro", "Section 2"]
    assert any(s.kind is SpanKind.CITATION for s in para.spans)
    masked, _ = mask(para.text, para.spans)
    assert masked.startswith("Moreover, the ⟦F1⟧ results")
    assert "⟦C1⟧" in masked and "⟦F2⟧" in masked
    assert [b.rewritable for b in doc.blocks] == [False, True, True, False, False, False]


def test_export_rewrites_in_place_and_preserves_inline_elements() -> None:
    data = make_docx()
    doc = parse(data, DocFormat.DOCX)
    para = doc.blocks[1]
    new = (
        "The ⟦F1⟧ results matter a great deal here ⟦C1⟧, as ⟦F2⟧ shows. "
        "Still, 30% of samples failed."
    )
    out = export(data, doc, {para.id: new})

    reparsed = parse(out, DocFormat.DOCX)
    assert reparsed.blocks[1].text == (
        "The in vitro results matter a great deal here (Smith, 2020), as Section 2 shows. "
        "Still, 30% of samples failed."
    )
    for before, after in zip(doc.blocks, reparsed.blocks, strict=True):
        if before.id != para.id:
            assert (before.text, before.kind) == (after.text, after.kind)

    body = docx.Document(io.BytesIO(out)).element.body
    p = next(p for i, p, _ in iter_paragraphs(body) if i == para.anchor["p_index"])
    italic = [
        "".join(t.text for t in r.iter(qn("w:t")))
        for r in p.iter(qn("w:r"))
        if r.find(qn("w:rPr")) is not None and r.find(qn("w:rPr")).find(qn("w:i")) is not None
    ]
    assert italic == ["in vitro"]
    assert p.find(qn("w:hyperlink")).get(qn("w:anchor")) == "sec2"
    tags = [c.tag for c in p]
    assert tags.index(qn("w:bookmarkStart")) < tags.index(qn("w:bookmarkEnd"))


def test_export_rejects_changed_document() -> None:
    data = make_docx()
    doc = parse(data, DocFormat.DOCX)
    doc.blocks[1] = doc.blocks[1].model_copy(update={"text": "something else"})
    with pytest.raises(ValueError, match="changed"):
        export(data, doc, {doc.blocks[1].id: "x"})


def test_export_leaves_every_other_package_part_byte_identical() -> None:
    import zipfile

    data = make_docx()
    doc = parse(data, DocFormat.DOCX)
    para = doc.blocks[2]
    out = export(data, doc, {para.id: "A rewritten bullet point that is still long enough here."})
    before, after = zipfile.ZipFile(io.BytesIO(data)), zipfile.ZipFile(io.BytesIO(out))
    assert before.namelist() == after.namelist()
    changed = [n for n in before.namelist() if before.read(n) != after.read(n)]
    assert changed == ["word/document.xml"]
