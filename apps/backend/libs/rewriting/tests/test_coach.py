import io
from typing import Any

import docx
import pytest
from docx.oxml.ns import qn
from pydantic import BaseModel

from detection import StylometricProvider, analyze
from docmodel import DocFormat, parse
from llm_client.types import ChatMessage
from rewriting import load_modes
from rewriting.coach import Edit, build_view, check_sentence, export_with_edits, suggest

TEXT = (
    "Moreover, it is important to note that 30% of pupils (Smith, 2020) may need support. "
    "My grandmother never trusted the microwave, and she said so every single day."
)


class JudgeLLM:
    def __init__(self, alternatives: list[dict[str, str]]) -> None:
        self.alternatives = alternatives
        self.calls: list[list[ChatMessage]] = []

    async def complete(self, messages: list[ChatMessage], **kw: Any) -> str:
        raise AssertionError("coach only uses judge()")

    async def judge[M: BaseModel](
        self, messages: list[ChatMessage], schema: type[M], *, model: str
    ) -> M:
        assert model == "rewrite"
        self.calls.append(messages)
        return schema.model_validate({"alternatives": self.alternatives})


@pytest.fixture
def mode() -> Any:
    return load_modes()["student.essay"]


def make_docx() -> bytes:
    d = docx.Document()
    d.add_heading("Introduction", level=1)
    p = d.add_paragraph()
    p.add_run("Moreover, it is important to note that 30% of ")
    p.add_run("pupils").italic = True
    p.add_run(" (Smith, 2020) may need support. My grandmother never trusted the microwave.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


async def test_view_marks_flagged_sentences() -> None:
    doc = parse(f"Intro\n\n{TEXT}\n".encode(), DocFormat.TXT)
    view = build_view(doc, await analyze(doc, [StylometricProvider("essay")]))
    para = next(b for b in view.blocks if b.kind == "paragraph")
    assert para.editable and para.sentences[0].band in ("elevated", "uncertain")
    assert any("stock phrase" in r for r in para.sentences[0].reasons)
    assert view.flagged >= 1


async def test_check_sentence_explains_problems(mode: Any) -> None:
    doc = parse(f"{TEXT}\n".encode(), DocFormat.TXT)
    block = doc.blocks[0]
    good = await check_sentence(mode, block, 0, "About 30% of pupils (Smith, 2020) may need help.")
    assert good.passed, good.checks
    bad = await check_sentence(mode, block, 0, "About 35% of pupils need help.")
    messages = {c.gate: c.message for c in bad.checks if not c.passed}
    assert "“(Smith, 2020)”" in messages["G1_placeholders"]
    assert "35" in messages["G2_numbers"]
    assert "hedge" in messages["G4_polarity_modality"]


async def test_suggestions_are_verified_and_sorted(mode: Any) -> None:
    doc = parse(f"{TEXT}\n".encode(), DocFormat.TXT)
    llm = JudgeLLM(
        [
            {"approach": "minimal", "text": "It is worth saying that 30% of pupils need support."},
            {"approach": "restructured", "text": "Some 30% of pupils ⟦C1⟧ may need support."},
            {"approach": "plain", "text": "Some 30% of pupils ⟦C1⟧ may need support."},
        ]
    )
    out = await suggest(llm, mode, doc, doc.blocks[0].id, 0, reasons=["stock phrase"])
    assert [s.approach for s in out] == ["restructured", "minimal"]  # duplicate dropped
    assert out[0].check.passed
    assert out[0].text == "Some 30% of pupils (Smith, 2020) may need support."
    assert not out[1].check.passed
    user = llm.calls[0][-1].text()
    assert "⟦C1⟧ = citation: '(Smith, 2020)'" in user and "stock phrase" in user


async def test_export_applies_edits_and_keeps_formatting(mode: Any) -> None:
    data = make_docx()
    doc = parse(data, DocFormat.DOCX)
    para = doc.blocks[1]
    assert para.rewritable
    edit = Edit(
        block_id=para.id, sentence_index=0, text="About 30% of pupils (Smith, 2020) may need help."
    )
    result = await export_with_edits(data, DocFormat.DOCX, doc, [edit], [StylometricProvider()])
    assert result.warnings == []
    out = docx.Document(io.BytesIO(result.output))
    p = out.paragraphs[1]
    assert p.text == (
        "About 30% of pupils (Smith, 2020) may need help. My grandmother never trusted the "
        "microwave."
    )
    italic = [r.text for r in p.runs if r._r.find(qn("w:rPr")) is not None and r.italic]
    assert italic == ["pupils"]

    dropped = Edit(block_id=para.id, sentence_index=0, text="Many pupils may need help.")
    result = await export_with_edits(data, DocFormat.DOCX, doc, [dropped], [StylometricProvider()])
    assert any("(Smith, 2020)" in w for w in result.warnings)


def test_voices_load_with_attribution_and_reach_prompts() -> None:
    from rewriting.prompts import system_prompt
    from rewriting.voices import load_voices, measure

    voices = load_voices()
    assert {"classic-essay", "academic-article"} <= set(voices)
    for v in voices.values():
        assert v.license and v.source and 3 <= len(v.exemplars) <= 6
        assert 10 < v.targets.sentence_words_mean < 40
    prompt = system_prompt(load_modes()["student.essay"], "moderate", voices["academic-article"])
    assert "Academic article" in prompt and "Formal, subject-specific vocabulary" in prompt
    assert "never reuse their facts" in prompt

    t = measure(["I can't say. We went home! Did we? It's late, and the rain kept falling."])
    assert t.contractions_per_100_words > 0 and t.first_person_per_100_words > 0
    assert t.question_share == 0.25
