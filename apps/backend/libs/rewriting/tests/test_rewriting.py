from collections.abc import Callable
from typing import Any

import pytest
from pydantic import BaseModel

from detection import StylometricProvider, analyze
from docmodel import DocFormat, export, parse
from llm_client import GatewayError
from llm_client.types import ChatMessage
from rewriting import EngineConfig, LlmEntailmentScorer, RewriteEngine, load_modes
from rewriting.prompts import clean_output

Reply = Callable[[str, float, int], str]


class FakeLLM:
    def __init__(self, reply: Reply, judgements: dict[str, Any] | None = None) -> None:
        self._reply = reply
        self.prompts: list[list[ChatMessage]] = []
        self.judge_calls: list[list[ChatMessage]] = []
        self._judgements = judgements

    async def complete(
        self, messages: list[ChatMessage], *, model: str, temperature: float, max_tokens: int
    ) -> str:
        assert model == "rewrite"
        self.prompts.append(messages)
        return self._reply(messages[-1].text(), temperature, len(self.prompts))

    async def judge[M: BaseModel](
        self, messages: list[ChatMessage], schema: type[M], *, model: str
    ) -> M:
        self.judge_calls.append(messages)
        return schema.model_validate(self._judgements or {"results": []})


TEXT = """Introduction

Moreover, it is important to note that technology plays a crucial role in education \
(Smith, 2020). Digital tools may improve outcomes for some students.

Appendix

Furthermore, it is important to note that this appendix plays a crucial role in the study.
"""

GOOD = "Technology shapes education ⟦C1⟧, and digital tools may improve outcomes for some students."
DROPS_CITATION = (
    "Technology shapes education, and digital tools may improve outcomes for some students."
)


@pytest.fixture
def modes() -> dict[str, Any]:
    return load_modes()


def test_modes_load(modes: dict[str, Any]) -> None:
    assert set(modes) == {
        "student.essay",
        "student.research_paper",
        "student.story",
        "professional.technical_docs",
        "professional.business_docs",
    }
    paper = modes["student.research_paper"]
    assert paper.freedom_for("2. Materials and Methods") == "minimal"
    assert paper.freedom_for("References") == "locked"
    assert paper.freedom_for("Discussion") == "conservative"
    assert paper.gate_policy().strict_hedges
    assert modes["professional.business_docs"].gate_policy().strict_deontic


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("Here is the rewritten paragraph:\nA new text.", "A new text."),
        ('"A quoted answer."', "A quoted answer."),
        ("```\nFenced text.\n```", "Fenced text."),
        ("<<<\nMarked\ntext.\n>>>", "Marked text."),
        ('He said "hi" and "bye".', 'He said "hi" and "bye".'),
    ],
)
def test_clean_output(raw: str, expected: str) -> None:
    assert clean_output(raw, "Original single line.") == expected


async def test_rewrite_passes_and_exports(modes: dict[str, Any]) -> None:
    doc = parse(TEXT.encode(), DocFormat.TXT)
    llm = FakeLLM(lambda prompt, t, n: GOOD)
    engine = RewriteEngine(llm, modes["student.essay"], config=EngineConfig(temperatures=[0.7]))
    block = doc.blocks[1]
    outcome = await engine.rewrite_block(block, reasons=["stock phrase: “plays a crucial role”"])
    assert outcome.status == "rewritten", outcome.model_dump()
    assert outcome.rewrite == (
        "Technology shapes education (Smith, 2020), and digital tools may improve outcomes "
        "for some students."
    )
    assert outcome.ai_score_after is not None and outcome.ai_score_before is not None
    assert outcome.ai_score_after < outcome.ai_score_before
    user = llm.prompts[0][-1].text()
    assert "⟦C1⟧ = citation: '(Smith, 2020)'" in user and "plays a crucial role" in user
    system = llm.prompts[0][0].text()
    assert "Mode: Essay (student)" in system

    out = export(TEXT.encode(), doc, {block.id: outcome.rewrite_masked or ""}).decode()
    assert outcome.rewrite in out and "Appendix" in out


async def test_retry_with_feedback_then_success(modes: dict[str, Any]) -> None:
    doc = parse(TEXT.encode(), DocFormat.TXT)
    llm = FakeLLM(lambda prompt, t, n: GOOD if "rejected" in prompt else DROPS_CITATION)
    engine = RewriteEngine(llm, modes["student.essay"], config=EngineConfig(temperatures=[0.7]))
    outcome = await engine.rewrite_block(doc.blocks[1])
    assert outcome.status == "rewritten"
    assert [c.round for c in outcome.candidates] == [1, 2]
    assert "G1_placeholders" in llm.prompts[1][-1].text()


async def test_never_returns_failing_rewrite(modes: dict[str, Any]) -> None:
    doc = parse(TEXT.encode(), DocFormat.TXT)
    hedge_removed = "Technology shapes education ⟦C1⟧, and digital tools improve outcomes."
    engine = RewriteEngine(FakeLLM(lambda p, t, n: hedge_removed), modes["student.essay"])
    outcome = await engine.rewrite_block(doc.blocks[1])
    assert outcome.status == "kept_original" and outcome.reason == "failed_verification"
    assert outcome.rewrite is None and len(outcome.candidates) == 4


async def test_llm_errors_keep_original(modes: dict[str, Any]) -> None:
    def boom(prompt: str, t: float, n: int) -> str:
        raise GatewayError(503, "all_deployments_failed", "no provider")

    doc = parse(TEXT.encode(), DocFormat.TXT)
    outcome = await RewriteEngine(FakeLLM(boom), modes["student.essay"]).rewrite_block(
        doc.blocks[1]
    )
    assert (outcome.status, outcome.reason) == ("kept_original", "llm_unavailable")


async def test_locked_section_is_skipped(modes: dict[str, Any]) -> None:
    doc = parse(TEXT.encode(), DocFormat.TXT)
    appendix = doc.blocks[3]
    assert appendix.section == "Appendix" and appendix.rewritable
    llm = FakeLLM(lambda p, t, n: GOOD)
    outcome = await RewriteEngine(llm, modes["student.research_paper"]).rewrite_block(appendix)
    assert (outcome.status, outcome.reason) == ("skipped", "section_locked")
    assert llm.prompts == []


async def test_entailment_gate_blocks_unsupported_rewrite(modes: dict[str, Any]) -> None:
    doc = parse(TEXT.encode(), DocFormat.TXT)
    judgements = {
        "results": [
            {"id": 0, "label": "entailed", "confidence": 0.95},
            {"id": 1, "label": "neutral", "confidence": 0.9},
            {"id": 2, "label": "entailed", "confidence": 0.95},
        ]
    }
    llm = FakeLLM(lambda p, t, n: GOOD, judgements)
    engine = RewriteEngine(
        llm,
        modes["student.essay"],
        scorer=LlmEntailmentScorer(llm),
        config=EngineConfig(temperatures=[0.7], max_rounds=1),
    )
    outcome = await engine.rewrite_block(doc.blocks[1])
    assert outcome.status == "kept_original"
    g5 = outcome.candidates[0].gates[-1]
    assert g5.gate == "G5_entailment" and not g5.passed
    assert "[ref]" in llm.judge_calls[0][-1].text()


async def test_document_selection_uses_analysis(modes: dict[str, Any]) -> None:
    doc = parse(TEXT.encode(), DocFormat.TXT)
    result = await analyze(doc, [StylometricProvider("essay")])
    llm = FakeLLM(lambda p, t, n: GOOD)
    engine = RewriteEngine(llm, modes["student.essay"], config=EngineConfig(temperatures=[0.7]))
    selected = engine.select_blocks(doc, result, None)
    assert [b.id for b in selected] == [
        b.id for b in doc.blocks if b.rewritable and b.section == "Introduction"
    ] + [doc.blocks[3].id]
    outcomes = await engine.rewrite_document(doc, block_ids={doc.blocks[1].id})
    assert [o.block_id for o in outcomes] == [doc.blocks[1].id]


async def test_pipeline_docx_changes_only_rewritten_paragraph(modes: dict[str, Any]) -> None:
    import io
    import zipfile

    import docx

    from detection import StylometricProvider
    from rewriting.pipeline import run_pipeline

    d = docx.Document()
    d.add_heading("Introduction", level=1)
    d.add_paragraph(
        "Moreover, it is important to note that technology plays a crucial role in education "
        "(Smith, 2020). Furthermore, digital tools unlock the full potential of learners."
    )
    d.add_paragraph("My grandmother never trusted the microwave, and she said so every single day.")
    buf = io.BytesIO()
    d.save(buf)
    data = buf.getvalue()

    tail = ", and digital tools help learners reach their full potential."
    rewrite = "Technology shapes education ⟦C1⟧" + tail
    stages: list[str] = []
    result = await run_pipeline(
        data,
        DocFormat.DOCX,
        modes["student.essay"],
        llm=FakeLLM(lambda p, t, n: rewrite),
        providers=[StylometricProvider("essay")],
        config=EngineConfig(temperatures=[0.7]),
        on_progress=lambda stage, done, total: stages.append(stage),
    )
    assert [p.status for p in result.paragraphs] == ["rewritten"]
    para = result.paragraphs[0]
    assert para.rewrite and "(Smith, 2020)" in para.rewrite
    assert para.score_before is not None and para.score_after is not None
    assert para.score_after < para.score_before
    assert result.before.score is not None and result.after.score is not None
    assert result.after.score < result.before.score
    assert stages[0] == "analyzing" and stages[-1] == "done" and "rewriting" in stages

    before, after = zipfile.ZipFile(io.BytesIO(data)), zipfile.ZipFile(io.BytesIO(result.output))
    assert [n for n in before.namelist() if before.read(n) != after.read(n)] == [
        "word/document.xml"
    ]
    texts = [p.text for p in docx.Document(io.BytesIO(result.output)).paragraphs]
    assert texts[0] == "Introduction" and texts[2].startswith("My grandmother")
    assert texts[1] == "Technology shapes education (Smith, 2020)" + tail
